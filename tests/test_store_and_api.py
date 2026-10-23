import subprocess
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from evoreport.api import create_app
from evoreport.config import Settings
from evoreport.store import Store


class StoreAndApiTests(unittest.TestCase):
    def test_daily_queue_status_full_and_timeout_responses(self):
        with tempfile.TemporaryDirectory() as directory:
            client = TestClient(create_app(Settings("127.0.0.1", 8090, Path(directory) / "api.db")))
            runtime = client.app.state.runtime
            payload = {"report_date": "2026-10-06", "completed": "完成联调"}
            with runtime.daily_queue.acquire(1):
                self.assertEqual(1, client.get("/api/agents/daily-queue").json()["running"])
                runtime.daily_queue.max_waiting = 0
                response = client.post("/api/reports/daily", json=payload)
                self.assertEqual(503, response.status_code)
                self.assertIn("队列已满", response.json()["detail"])
                runtime.daily_queue.max_waiting = 5
                runtime.MAX_RUN_SECONDS = 0.02
                response = client.post("/api/reports/daily", json=payload)
                self.assertEqual(504, response.status_code)
                self.assertIn("等待超时", response.json()["detail"])
            self.assertEqual(0, client.get("/api/agents/daily-queue").json()["waiting"])
            self.assertEqual(2, len(client.app.state.store.list_reports()))
            self.assertEqual([], client.app.state.store.list_agent_runs())

    def test_confirmed_report_becomes_evaluation_case(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            report = store.create_report(
                "daily", "2026-10-06", {"kind": "daily", "date": "2026-10-06", "facts": []}, "draft", 1
            )
            confirmed = store.confirm_report(report["id"], "# final")
            cases = store.list_evaluation_cases()
            self.assertEqual("confirmed", confirmed["status"])
            self.assertEqual(1, len(cases))
            self.assertEqual("# final", cases[0]["expected_text"])

    def test_draft_save_does_not_confirm_or_overwrite_final_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            report = client.post("/api/reports/daily", json={
                "report_date": "2026-10-06", "completed": "完成联调",
            }).json()
            saved = client.patch(f"/api/reports/{report['id']}/draft", json={"draft": "# 修改后的草稿"})
            self.assertEqual(200, saved.status_code)
            self.assertEqual("# 修改后的草稿", saved.json()["draft"])
            self.assertEqual("draft", saved.json()["status"])
            self.assertEqual([], client.app.state.store.list_evaluation_cases())

            review = client.post(f"/api/reports/{report['id']}/confirm", json={"final_text": "# 确认终稿"})
            self.assertEqual(409, review.status_code)
            self.assertFalse(review.json()["detail"]["review"]["unsafe"])
            confirmed = client.post(f"/api/reports/{report['id']}/confirm", json={
                "final_text": "# 确认终稿", "acknowledge_warnings": True,
            })
            self.assertEqual(200, confirmed.status_code)
            rejected = client.patch(f"/api/reports/{report['id']}/draft", json={"draft": "误写终稿"})
            self.assertEqual(409, rejected.status_code)
            current = client.get(f"/api/reports/{report['id']}").json()
            self.assertEqual("# 修改后的草稿", current["draft"])
            self.assertEqual("# 确认终稿", current["final_text"])
            self.assertEqual(1, len(client.app.state.store.list_evaluation_cases()))

    def test_api_adds_project_and_generates_daily_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            repository = root / "repo"
            repository.mkdir()
            self._git(repository, "init")
            self._git(repository, "config", "user.name", "EvoReport Test")
            self._git(repository, "config", "user.email", "test@example.invalid")
            (repository / "app.py").write_text("print('ok')\n", encoding="utf-8")
            self._git(repository, "add", "app.py")
            self._git(repository, "commit", "-m", "create application")
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))

            created = client.post("/api/projects", json={"name": "Demo", "path": str(repository)})
            self.assertEqual(201, created.status_code)
            report = client.post("/api/reports/daily", json={
                "report_date": date.today().isoformat(),
                "project_ids": [created.json()["id"]],
                "completed": "完成页面联调",
                "problems": "",
                "tomorrow": "补充测试",
            })
            self.assertEqual(201, report.status_code)
            self.assertIn("Demo：create application", report.json()["draft"])
            self.assertEqual("git_commit", report.json()["source"]["fact_sources"][1]["source_type"])
            self.assertEqual("Demo", report.json()["source"]["fact_sources"][1]["project_name"])
            self.assertEqual(created.json()["id"], report.json()["source"]["fact_sources"][1]["project_id"])
            self.assertEqual(40, len(report.json()["source"]["fact_sources"][1]["commit_hash"]))
            trace_steps = report.json()["agent_trace"]["steps"]
            self.assertEqual("Planner Agent", trace_steps[0]["agent_name"])
            planned_tasks = trace_steps[0]["output"]["task_ids"]
            self.assertEqual(planned_tasks, [step["input"]["task_id"] for step in trace_steps
                                             if step["agent_name"] == "Fact Agent"])
            self.assertEqual(["Writer Agent", "Critic Agent", "Gate Agent"],
                             [step["agent_name"] for step in trace_steps[-3:]])
            self.assertEqual("Gate Agent", report.json()["agent_trace"]["steps"][-1]["agent_name"])

    def test_api_adds_plain_project_and_generates_manual_daily_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            plain_project = root / "plain-project"
            plain_project.mkdir()
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))

            created = client.post("/api/projects", json={
                "name": "课程作业", "path": str(plain_project), "project_type": "plain"
            })
            self.assertEqual(201, created.status_code)
            self.assertEqual("plain", created.json()["project_type"])
            report = client.post("/api/reports/daily", json={
                "report_date": date.today().isoformat(),
                "project_ids": [created.json()["id"]],
                "completed": "完成算法作业", "problems": "", "tomorrow": "整理注释",
            })
            self.assertEqual(201, report.status_code)
            self.assertIn("完成算法作业", report.json()["draft"])
            self.assertNotIn("开发统计", report.json()["draft"])

    def test_report_archive_filters_by_snapshotted_project_and_date(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name in ("one", "two"):
                (root / name).mkdir()
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "archive.db")))
            projects = [client.post("/api/projects", json={
                "name": name, "path": str(root / name), "project_type": "plain",
            }).json() for name in ("one", "two")]

            def daily(day, project_ids, completed):
                response = client.post("/api/reports/daily", json={
                    "report_date": day, "project_ids": project_ids, "completed": completed,
                })
                self.assertEqual(201, response.status_code, response.text)
                return response.json()

            first = daily("2026-10-01", [projects[0]["id"]], "完成第一项")
            second = daily("2026-10-02", [projects[1]["id"]], "完成第二项")
            shared = daily("2026-10-03", [projects[0]["id"], projects[1]["id"]], "共同完成联调")
            unassigned = daily("2026-10-03", [], "独立记录")
            old = client.app.state.store.create_report(
                "daily", "2026-10-04", {"kind": "daily", "facts": ["旧事实"]}, "旧草稿", 1,
            )
            self.assertEqual([], first["source"]["projects"][0]["activity"]["commits"])
            self.assertEqual("manual", shared["source"]["fact_sources"][0]["source_type"])
            self.assertEqual(
                {first["id"], shared["id"]},
                {item["id"] for item in client.get(f"/api/reports?project_id={projects[0]['id']}").json()},
            )
            self.assertEqual(
                {second["id"], shared["id"]},
                {item["id"] for item in client.get(f"/api/reports?project_id={projects[1]['id']}").json()},
            )
            narrowed = client.get(f"/api/reports?project_id={projects[0]['id']}&start=2026-10-03&end=2026-10-03")
            self.assertEqual([shared["id"]], [item["id"] for item in narrowed.json()])
            self.assertEqual(
                {first["id"], second["id"], shared["id"], unassigned["id"], old["id"]},
                {item["id"] for item in client.get("/api/reports").json()},
            )
            self.assertEqual(400, client.get("/api/reports?start=2026-10-04&end=2026-10-03").status_code)
            self.assertEqual(422, client.get("/api/reports?start=bad-date").status_code)

    def test_quick_daily_collects_today_notes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            project = root / "plain-project"
            project.mkdir()
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            created = client.post("/api/projects", json={
                "name": "普通项目", "path": str(project), "project_type": "plain"
            }).json()

            note = client.post("/api/notes", json={
                "note_date": date.today().isoformat(), "project_id": created["id"], "content": "解决了启动报错"
            })
            self.assertEqual(201, note.status_code)
            report = client.post("/api/reports/daily/quick", json={"report_date": date.today().isoformat()})
            self.assertEqual(201, report.status_code)
            self.assertIn("解决了启动报错", report.json()["draft"])
            self.assertEqual(1, len(report.json()["source"]["quick_notes"]))

    def test_import_historical_report_creates_confirmed_evaluation_sample(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            response = client.post("/api/reports/historical", json={
                "report_date": "2026-07-01",
                "content": "# 日报\n\n1. 完成新闻内容更新\n2. 修复移动端导航问题",
            })
            self.assertEqual(201, response.status_code)
            imported = response.json()
            self.assertEqual("historical", imported["origin"])
            self.assertEqual("confirmed", imported["status"])
            self.assertTrue(any("修复移动端" in item for item in imported["source"]["facts"]))
            self.assertEqual(1, len(client.app.state.store.list_evaluation_cases()))

    def test_synthetic_samples_are_not_evaluation_cases(self):
        with tempfile.TemporaryDirectory() as directory:
            settings = Settings("127.0.0.1", 8090, Path(directory) / "api.db")
            client = TestClient(create_app(settings))
            response = client.post("/api/reports/synthetic-samples")
            self.assertEqual(201, response.status_code)
            self.assertEqual(8, response.json()["created"])
            reports = client.get("/api/reports").json()
            self.assertEqual(8, sum(report["origin"] == "synthetic" for report in reports))
            self.assertEqual([], client.app.state.store.list_evaluation_cases())

    def test_exports_confirmed_reports_and_creates_sqlite_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            imported = client.post("/api/reports/historical", json={
                "report_date": "2026-07-01", "content": "# 日报\n\n## 今日完成\n\n1. 完成接口联调",
            }).json()
            single = client.get(f"/api/reports/{imported['id']}/export")
            archive = client.get("/api/reports/export/confirmed")
            backup = client.post("/api/backups")
            self.assertEqual(200, single.status_code)
            self.assertEqual(200, archive.status_code)
            self.assertIn("完成接口联调", archive.text)
            self.assertEqual(201, backup.status_code)
            self.assertTrue(backup.json()["filename"].endswith(".db"))

    def test_word_export_preserves_daily_content_and_lists(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            imported = client.post("/api/reports/historical", json={
                "report_date": "2026-07-01",
                "content": "# 工作日报 · 2026-07-01\n\n## 今日完成\n1. 修复**移动端导航**\n- 完成联调\n\n## 明日计划\n继续测试",
            }).json()
            response = client.get(f"/api/reports/{imported['id']}/export/docx")
            self.assertEqual(200, response.status_code)
            self.assertTrue(response.content.startswith(b"PK"))
            document = Document(BytesIO(response.content))
            paragraphs = [item.text for item in document.paragraphs]
            self.assertEqual("Title", document.paragraphs[0].style.name)
            self.assertIn("2026年7月1日工作日报", paragraphs)
            self.assertIn("今日完成", paragraphs)
            self.assertIn("1. 修复移动端导航", paragraphs)
            self.assertIn("• 完成联调", paragraphs)
            self.assertIn("继续测试", paragraphs)
            self.assertEqual(404, client.get("/api/reports/missing/export/docx").status_code)

            descriptive = client.post("/api/reports/historical", json={
                "report_date": "2026-07-02", "content": "# 工作日报 Thyseed 俄罗斯官网\n\n完成新闻页修改",
            }).json()
            descriptive_docx = client.get(f"/api/reports/{descriptive['id']}/export/docx")
            self.assertIn(
                "工作日报 Thyseed 俄罗斯官网",
                [item.text for item in Document(BytesIO(descriptive_docx.content)).paragraphs],
            )

    def test_preview_and_restore_keep_a_safety_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            first = client.post("/api/reports/historical", json={
                "report_date": "2026-07-01", "content": "# 第一篇真实日报",
            }).json()
            backup = client.post("/api/backups").json()
            second = client.post("/api/reports/historical", json={
                "report_date": "2026-07-02", "content": "# 第二篇真实日报",
            }).json()
            backup_bytes = (root / "backups" / backup["filename"]).read_bytes()
            preview = client.post(
                "/api/backups/preview", content=backup_bytes,
                headers={"Content-Type": "application/octet-stream"},
            )
            self.assertEqual(200, preview.status_code)
            self.assertEqual(1, preview.json()["counts"]["reports"])
            restored = client.post("/api/backups/restore", json={
                "preview_id": preview.json()["preview_id"], "confirmation": "RESTORE",
            })
            self.assertEqual(200, restored.status_code)
            self.assertEqual([first["id"]], [report["id"] for report in client.get("/api/reports").json()])
            safety = root / "backups" / restored.json()["safety_backup"]
            self.assertEqual(2, Store.inspect_backup(safety)["reports"])
            self.assertNotEqual(first["id"], second["id"])

    def test_invalid_backup_is_rejected_without_changing_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            client.post("/api/reports/historical", json={
                "report_date": "2026-07-01", "content": "# 保留的日报",
            })
            preview = client.post(
                "/api/backups/preview", content=b"not sqlite",
                headers={"Content-Type": "application/octet-stream"},
            )
            self.assertEqual(400, preview.status_code)
            self.assertEqual(1, len(client.get("/api/reports").json()))

    def test_parses_natural_language_summary_without_saving_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            client = TestClient(create_app(Settings("127.0.0.1", 8090, root / "api.db")))
            response = client.post("/api/reports/parse-work-summary", json={
                "content": "今日完成修复对账异常并完成接口联调。遇到的问题是原始时间缺失。明日计划继续优化导出功能。"
            })
            self.assertEqual(200, response.status_code)
            self.assertIn("修复对账异常", response.json()["completed"])
            self.assertIn("原始时间缺失", response.json()["problems"])
            self.assertIn("优化导出功能", response.json()["tomorrow"])
            self.assertEqual([], client.get("/api/reports").json())

            labeled = client.post("/api/reports/parse-work-summary", json={
                "content": "今日完成：接口联调；遇到问题：排查后调整字段映射；明日计划：继续测试导出。"
            })
            self.assertEqual(200, labeled.status_code)
            self.assertEqual("接口联调", labeled.json()["completed"])
            self.assertEqual("排查后调整字段映射", labeled.json()["problems"])
            self.assertEqual("继续测试导出", labeled.json()["tomorrow"])
            self.assertEqual([], client.get("/api/reports").json())

    def test_documented_sample_parse_does_not_create_report_or_case(self):
        sample_path = Path(__file__).resolve().parents[1] / "测试用随手记样例.md"
        paragraph = next(
            line for line in sample_path.read_text(encoding="utf-8").splitlines()
            if line.startswith("今日完成：完成接口联调并整理测试用例")
        )
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "api.db"
            client = TestClient(create_app(Settings("127.0.0.1", 8090, database_path)))
            response = client.post("/api/reports/parse-work-summary", json={"content": paragraph})
            self.assertEqual(200, response.status_code)
            self.assertIn("排查、调整", response.json()["problems"])
            self.assertEqual([], client.get("/api/reports").json())
            self.assertEqual([], Store(database_path).list_evaluation_cases())

    @staticmethod
    def _git(path: Path, *arguments: str) -> None:
        completed = subprocess.run(["git", "-C", str(path), *arguments], capture_output=True, text=True, encoding="utf-8")
        if completed.returncode != 0:
            raise AssertionError(completed.stderr)


if __name__ == "__main__":
    unittest.main()
