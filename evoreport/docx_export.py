from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def export_report_docx(report: Dict[str, Any], destination: Path) -> None:
    """将一篇日报或周报的最终正文转换为可编辑的 Word 文档。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    _set_font(normal, 11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after in (
        ("Title", 18, 0, 14),
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 5),
    ):
        style = document.styles[name]
        _set_font(style, size)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    report_date = date.fromisoformat(report["report_date"])
    kind = "周报" if report["kind"] == "weekly" else "日报"
    title = f"{report_date.year}年{report_date.month}月{report_date.day}日工作{kind}"
    document.add_paragraph(title, style="Title")

    body = report["final_text"] or report["draft"]
    for raw_line in body.replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            text = heading.group(2).strip()
            if heading.group(1) == "#" and _is_duplicate_title(text, kind, report_date):
                continue
            paragraph = document.add_paragraph(style="Heading 1" if len(heading.group(1)) <= 2 else "Heading 2")
            _add_inline(paragraph, text)
            continue
        list_item = re.match(r"^(?:(\d+[.)、])|([-*]))\s+(.+)$", line)
        if list_item:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.27)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            prefix = list_item.group(1) or "•"
            paragraph.add_run(f"{prefix} ")
            _add_inline(paragraph, list_item.group(3))
            continue
        paragraph = document.add_paragraph()
        _add_inline(paragraph, line)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def _is_duplicate_title(text: str, kind: str, report_date: date) -> bool:
    return text in {kind, f"工作{kind}", f"工作{kind} · {report_date.isoformat()}"}


def _set_font(style: Any, size: int) -> None:
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_inline(paragraph: Any, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)
